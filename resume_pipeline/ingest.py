"""Turn a resume file on disk into provider-neutral content parts.

PDFs and images are passed through as raw bytes rather than run through a local
text extractor - the model reads layout, columns and tables far better than a
naive text dump does. Each provider client converts these parts into its own
wire format, so nothing here is Gemini- or Anthropic-specific.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import List, Union

MAX_REQUEST_BYTES = 18 * 1024 * 1024  # Inline request payloads must stay modest.

IMAGE_MEDIA_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
}
TEXT_SUFFIXES = {".txt", ".md", ".rst", ".text"}


@dataclass(frozen=True)
class TextPart:
    """A chunk of plain text."""

    text: str


@dataclass(frozen=True)
class BinaryPart:
    """A document or image, carried as raw bytes plus its MIME type."""

    data: bytes
    mime_type: str


ContentPart = Union[TextPart, BinaryPart]


class IngestError(ValueError):
    """Raised when a resume file cannot be read or is an unsupported type."""


def _docx_to_text(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise IngestError(
            "Reading .docx requires python-docx. Install it with: pip install python-docx"
        ) from exc

    document = docx.Document(str(path))
    parts: List[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise IngestError(f"No readable text found in {path.name}.")
    return text


def load_resume(path: str | Path) -> List[ContentPart]:
    """Return the provider-neutral content parts representing this resume."""
    path = Path(path)
    if not path.is_file():
        raise IngestError(f"Resume not found: {path}")

    size = path.stat().st_size
    if size == 0:
        raise IngestError(f"{path.name} is empty.")
    if size > MAX_REQUEST_BYTES:
        raise IngestError(
            f"{path.name} is {size / 1e6:.1f} MB, over the ~18 MB inline limit."
        )

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return [BinaryPart(data=path.read_bytes(), mime_type="application/pdf")]

    if suffix in IMAGE_MEDIA_TYPES:
        return [BinaryPart(data=path.read_bytes(), mime_type=IMAGE_MEDIA_TYPES[suffix])]

    if suffix == ".docx":
        return [TextPart(text=_docx_to_text(path))]

    if suffix == ".doc":
        raise IngestError(
            "Legacy .doc is not supported. Save it as .docx or .pdf and retry."
        )

    if suffix in TEXT_SUFFIXES or suffix == "":
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            raise IngestError(f"No readable text found in {path.name}.")
        return [TextPart(text=text)]

    raise IngestError(
        f"Unsupported resume format '{suffix}'. "
        "Supported: .pdf, .docx, .txt, .md, .png, .jpg, .webp, .gif"
    )


def load_job_description(path: str | Path | None) -> str:
    """Read an optional job-description file as plain text."""
    if not path:
        return ""
    path = Path(path)
    if not path.is_file():
        raise IngestError(f"Job description not found: {path}")
    if path.suffix.lower() == ".docx":
        return _docx_to_text(path)
    if path.suffix.lower() == ".pdf":
        raise IngestError("Job description must be a text file (.txt/.md/.docx), not a PDF.")
    return path.read_text(encoding="utf-8", errors="replace").strip()
